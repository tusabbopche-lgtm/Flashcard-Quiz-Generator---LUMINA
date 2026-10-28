import io
from typing import Tuple

from pypdf import PdfReader
from docx import Document


class ExtractionError(Exception):
    pass


def extract_text_from_file(filename: str, content: bytes) -> str:
    name_lower = filename.lower()
    if name_lower.endswith('.txt'):
        # Try utf-8, fallback to cp1252
        try:
            return content.decode('utf-8')
        except UnicodeError:
            return content.decode('cp1252', errors='replace')
    elif name_lower.endswith('.pdf'):
        # Check signature
        if not content.startswith(b'%PDF'):
            raise ExtractionError('Invalid PDF file')
        try:
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ExtractionError('Encrypted PDF')
            text_parts = []
            for page in reader.pages:
                try:
                    text_parts.append(page.extract_text() or '')
                except Exception:
                    continue
            full_text = '\n'.join(text_parts)
            if len(reader.pages) > 0:
                avg_chars = len(full_text) / len(reader.pages)
                if avg_chars < 50:
                    raise ExtractionError('Probable scanned PDF, paste the text instead')
            return full_text
        except ExtractionError:
            raise
        except Exception as e:
            raise ExtractionError(f'Failed to extract PDF: {e}')
    elif name_lower.endswith('.docx'):
        # Check it's a zip
        if not content.startswith(b'PK'):
            raise ExtractionError('Invalid DOCX file')
        try:
            doc = Document(io.BytesIO(content))
            parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            if p.text.strip():
                                parts.append(p.text)
            return '\n'.join(parts)
        except ExtractionError:
            raise
        except Exception as e:
            raise ExtractionError(f'Failed to extract DOCX: {e}')
    else:
        raise ExtractionError('Unsupported file type')
